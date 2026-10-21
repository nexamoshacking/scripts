import json
import os
import shutil
import sqlite3
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

#NXMS
CHROME_PROFILE_DIR = os.path.expanduser(
    r"~\AppData\Local\Google\Chrome\User Data\Default"
)
HISTORY_DB_PATH = os.path.join(CHROME_PROFILE_DIR, "History")
BOOKMARKS_PATH = os.path.join(CHROME_PROFILE_DIR, "Bookmarks")

TEMP_DB = "temp_chrome_history.sqlite"
OUTPUT_EXCEL = "Relatorio_Forense_Chrome.xlsx"


def get_history_and_downloads():
  """Extrai histórico de navegação e downloads do banco SQLite do Chrome."""
  if not os.path.exists(HISTORY_DB_PATH):
    print(f"[-] Arquivo de histórico não encontrado em: {HISTORY_DB_PATH}")
    return [], []

  # Copia para evitar bloqueio caso o Chrome esteja aberto
  shutil.copy2(HISTORY_DB_PATH, TEMP_DB)

  conn = sqlite3.connect(TEMP_DB)
  cursor = conn.cursor()

  # Query 1: Histórico
  query_history = """
    SELECT 
        urls.url, 
        urls.title, 
        urls.visit_count, 
        datetime(urls.last_visit_time/1000000 - 11644473600, 'unixepoch', 'localtime') AS visit_time
    FROM urls 
    WHERE urls.last_visit_time > 0
    ORDER BY urls.last_visit_time DESC 
    LIMIT 300
    """
  cursor.execute(query_history)
  history_data = cursor.fetchall()

  # Query 2: Downloads
  query_downloads = """
    SELECT 
        target_path,
        tab_url,
        total_bytes,
        datetime(start_time/1000000 - 11644473600, 'unixepoch', 'localtime') AS download_time
    FROM downloads
    ORDER BY start_time DESC
    LIMIT 100
    """
  cursor.execute(query_downloads)
  download_data = cursor.fetchall()

  conn.close()

  if os.path.exists(TEMP_DB):
    os.remove(TEMP_DB)

  return history_data, download_data


def get_bookmarks():
  """Lê e processa a árvore de favoritos (JSON) do Chrome."""
  if not os.path.exists(BOOKMARKS_PATH):
    print(f"[-] Arquivo de favoritos não encontrado em: {BOOKMARKS_PATH}")
    return []

  bookmarks_list = []

  try:
    with open(BOOKMARKS_PATH, "r", encoding="utf-8") as f:
      data = json.load(f)

    def parse_node(node, folder_path=""):
      node_type = node.get("type")
      name = node.get("name", "")

      if node_type == "folder":
        current_path = f"{folder_path}/{name}" if folder_path else name
        children = node.get("children", [])
        for child in children:
          parse_node(child, current_path)
      elif node_type == "url":
        url = node.get("url", "")
        # Chrome armazena date_added em microssegundos no formato WebKit Epoch (desde 1601)
        date_added_raw = int(node.get("date_added", 0))
        date_added_str = "N/A"
        if date_added_raw > 0:
          # Conversão rápida de timestamp WebKit para legível
          unix_ts = (date_added_raw / 1000000) - 11644473600
          if unix_ts > 0:
            import datetime

            date_added_str = datetime.datetime.fromtimestamp(
                unix_ts
            ).strftime("%Y-%m-%d %H:%M:%S")

        bookmarks_list.append([name, url, folder_path, date_added_str])

    roots = data.get("roots", {})
    for root_key in ["bookmark_bar", "other", "synced"]:
      if root_key in roots:
        parse_node(roots[root_key], folder_path=root_key)

  except Exception as e:
    print(f"[-] Erro ao ler favoritos: {e}")

  return bookmarks_list


def export_to_styled_xlsx(history_data, download_data, bookmarks_data):
  """Gera a planilha formatada em Excel."""
  wb = openpyxl.Workbook()

  # Estilos visuais
  header_fill = PatternFill(
      start_color="1F4E78", end_color="1F4E78", fill_type="solid"
  )
  header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
  data_font = Font(name="Calibri", size=10)

  thin_border = Border(
      left=Side(style="thin", color="D9D9D9"),
      right=Side(style="thin", color="D9D9D9"),
      top=Side(style="thin", color="D9D9D9"),
      bottom=Side(style="thin", color="D9D9D9"),
  )

  zebra_fill = PatternFill(
      start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"
  )

  # ABA 1: HISTÓRICO DE NAVEGAÇÃO
  ws_history = wb.active
  ws_history.title = "Histórico de Navegação"
  headers_history = [
      "URL Visita",
      "Título da Página",
      "Qtd. Visitas",
      "Data/Hora Local",
  ]
  ws_history.append(headers_history)
  for row in history_data:
    ws_history.append(list(row))
  format_sheet(
      ws_history,
      headers_history,
      header_fill,
      header_font,
      data_font,
      thin_border,
      zebra_fill,
  )

  # ABA 2: DOWNLOADS
  ws_downloads = wb.create_sheet(title="Downloads Realizados")
  headers_downloads = [
      "Caminho do Arquivo Salvo",
      "URL de Origem (Fonte)",
      "Tamanho (Bytes)",
      "Data/Hora Download",
  ]
  ws_downloads.append(headers_downloads)
  for row in download_data:
    ws_downloads.append(list(row))
  format_sheet(
      ws_downloads,
      headers_downloads,
      header_fill,
      header_font,
      data_font,
      thin_border,
      zebra_fill,
  )

  # ABA 3: FAVORITOS
  ws_bookmarks = wb.create_sheet(title="Favoritos Salvos")
  headers_bookmarks = [
      "Nome do Favorito",
      "URL",
      "Pasta / Caminho",
      "Data de Adição",
  ]
  ws_bookmarks.append(headers_bookmarks)
  for row in bookmarks_data:
    ws_bookmarks.append(row)
  format_sheet(
      ws_bookmarks,
      headers_bookmarks,
      header_fill,
      header_font,
      data_font,
      thin_border,
      zebra_fill,
  )

  wb.save(OUTPUT_EXCEL)
  print(f"[+] Relatório gerado com sucesso: {OUTPUT_EXCEL}")


def format_sheet(
    ws, headers, header_fill, header_font, data_font, border, zebra_fill
):
  """Aplica formatação às abas."""
  ws.views.sheetView[0].showGridLines = True

  for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )

  ws.freeze_panes = "A2"
  ws.auto_filter.ref = ws.dimensions

  for row_idx in range(2, ws.max_row + 1):
    is_even = row_idx % 2 == 0
    for col_idx in range(1, len(headers) + 1):
      cell = ws.cell(row=row_idx, column=col_idx)
      cell.font = data_font
      cell.border = border
      cell.alignment = Alignment(vertical="center")
      if is_even:
        cell.fill = zebra_fill

  for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
      val_str = str(cell.value or "")
      if len(val_str) > max_len:
        max_len = len(val_str)

    adjusted_width = min(max(max_len + 3, 15), 60)
    ws.column_dimensions[col_letter].width = adjusted_width


if __name__ == "__main__":
  print("[*] Extraindo artefatos do Chrome na máquina local...")
  history, downloads = get_history_and_downloads()
  bookmarks = get_bookmarks()

  print(
      f"[+] Registros encontrados: {len(history)} Histórico | {len(downloads)}"
      f" Downloads | {len(bookmarks)} Favoritos."
  )
  export_to_styled_xlsx(history, downloads, bookmarks)